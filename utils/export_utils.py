# =============================================================================
# FILE: utils/export_utils.py
# DESCRIPTION: CSV export utilities for data export functionality
# LOCATION: utils/export_utils.py
# PURPOSE: Centralized CSV generation for users, attendance, and statistics
# =============================================================================

"""
CSV export utilities for the Telegram School Management Bot.
"""

import csv
import io
from datetime import datetime, date
from typing import List, Dict, Any, Optional

from database.operations import (
    get_all_users,
    get_users_by_role,
    get_users_by_class,
    get_all_attendance_records,
    get_attendance_stats_by_class,
)
from database import get_db, Attendance, User, Class


def generate_users_csv(users: List[User]) -> str:
    """
    Generate CSV content for users.
    
    Args:
        users: List of User objects
        
    Returns:
        CSV content as string
    """
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Header
    writer.writerow([
        'ID',
        'Telegram ID',
        'Name',
        'Role',
        'Class ID',
        'Phone',
        'Birthday',
        'Language',
        'Created At'
    ])
    
    # Data rows
    role_map = {1: 'Student', 2: 'Teacher', 3: 'Leader', 4: 'Manager', 5: 'Developer'}
    
    for user in users:
        writer.writerow([
            user.id,
            user.telegram_id,
            user.name,
            role_map.get(user.role, 'Unknown'),
            user.class_id or '',
            user.phone or '',
            user.birthday.strftime('%Y-%m-%d') if user.birthday else '',
            user.language_preference or 'ar',
            user.created_at.strftime('%Y-%m-%d %H:%M:%S') if user.created_at else ''
        ])
    
    return output.getvalue()


def generate_attendance_csv(records: Optional[List[tuple]] = None) -> str:
    """
    Generate CSV content for attendance records.
    
    Args:
        records: Optional list of attendance records (if None, fetches all)
        
    Returns:
        CSV content as string
    """
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Header
    writer.writerow([
        'ID',
        'Student Name',
        'Student ID',
        'Class ID',
        'Date',
        'Status',
        'Note',
        'Marked By',
        'Marked At'
    ])
    
    # Fetch records if not provided
    if records is None:
        with get_db() as db:
            records = db.query(
                Attendance,
                User.name.label('student_name'),
                User.telegram_id.label('student_telegram_id')
            ).join(
                User, Attendance.user_id == User.id
            ).all()
    
    # Data rows
    for record in records:
        if hasattr(record, '_asdict'):
            # Named tuple from query
            att = record.Attendance if hasattr(record, 'Attendance') else record
            student_name = record.student_name if hasattr(record, 'student_name') else 'Unknown'
            student_tid = record.student_telegram_id if hasattr(record, 'student_telegram_id') else ''
        else:
            # Direct Attendance object
            att = record
            student_name = att.user.name if att.user else 'Unknown'
            student_tid = att.user.telegram_id if att.user else ''
        
        # Get marker name
        with get_db() as db:
            marker = db.query(User).filter(User.id == att.marked_by).first()
            marker_name = marker.name if marker else 'Unknown'
        
        writer.writerow([
            att.id,
            student_name,
            student_tid,
            att.class_id or '',
            att.date.strftime('%Y-%m-%d') if isinstance(att.date, date) else str(att.date),
            'Present' if att.status else 'Absent',
            att.note or '',
            marker_name,
            att.created_at.strftime('%Y-%m-%d %H:%M:%S') if att.created_at else ''
        ])
    
    return output.getvalue()


def generate_class_stats_csv() -> str:
    """
    Generate CSV content for class statistics.
    
    Returns:
        CSV content as string
    """
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Header
    writer.writerow([
        'Class ID',
        'Total Students',
        'Total Absences',
        'Absences with Reason',
        'Reason Rate (%)',
        'Top Reason 1',
        'Top Reason 2',
        'Top Reason 3'
    ])
    
    # Get all classes
    with get_db() as db:
        classes = db.query(Class).all()
    
    for class_obj in classes:
        # Get users and stats for this class
        users = get_users_by_class(class_obj.id)
        students = [u for u in users if u.role == 1]
        stats = get_attendance_stats_by_class(class_obj.id)
        
        total_absent = stats.get('total_absent', 0)
        with_reason = stats.get('total_with_reason', 0)
        reason_rate = (with_reason / total_absent * 100) if total_absent > 0 else 0
        
        # Get top 3 reasons
        reasons = stats.get('reason_breakdown', {})
        sorted_reasons = sorted(reasons.items(), key=lambda x: x[1], reverse=True)[:3]
        
        top_reasons = [r[0] for r in sorted_reasons] + ['', '', '']  # Pad with empty strings
        
        writer.writerow([
            class_obj.id,
            len(students),
            total_absent,
            with_reason,
            f'{reason_rate:.1f}',
            top_reasons[0],
            top_reasons[1],
            top_reasons[2]
        ])
    
    return output.getvalue()


def generate_full_report_csv() -> str:
    """
    Generate comprehensive CSV report combining users and attendance.
    
    Returns:
        CSV content as string
    """
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Header
    writer.writerow([
        'Student ID',
        'Student Name',
        'Telegram ID',
        'Class ID',
        'Phone',
        'Total Attendance Records',
        'Present Count',
        'Absent Count',
        'Attendance Rate (%)',
        'Latest Attendance Date',
        'Latest Status'
    ])
    
    # Get all students
    students = get_users_by_role(1)  # Role 1 = Student
    
    for student in students:
        # Get attendance stats for this student
        with get_db() as db:
            records = db.query(Attendance).filter(Attendance.user_id == student.id).all()
            
            total_records = len(records)
            present_count = sum(1 for r in records if r.status)
            absent_count = total_records - present_count
            attendance_rate = (present_count / total_records * 100) if total_records > 0 else 0
            
            # Get latest record
            latest = db.query(Attendance).filter(
                Attendance.user_id == student.id
            ).order_by(Attendance.date.desc()).first()
            
            latest_date = latest.date.strftime('%Y-%m-%d') if latest else ''
            latest_status = 'Present' if (latest and latest.status) else ('Absent' if latest else '')
        
        writer.writerow([
            student.id,
            student.name,
            student.telegram_id,
            student.class_id or '',
            student.phone or '',
            total_records,
            present_count,
            absent_count,
            f'{attendance_rate:.1f}',
            latest_date,
            latest_status
        ])
    
    return output.getvalue()


def save_csv_to_file(csv_content: str, filename: str) -> str:
    """
    Save CSV content to a file.
    
    Args:
        csv_content: CSV content as string
        filename: Filename (without path)
        
    Returns:
        Full file path
    """
    import os
    
    # Create exports directory if it doesn't exist
    export_dir = "/home/Bisho/Telegram/exports"
    if not os.path.exists(export_dir):
        os.makedirs(export_dir)
    
    # Generate filename with timestamp
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    full_filename = f"{filename}_{timestamp}.csv"
    file_path = os.path.join(export_dir, full_filename)
    
    # Write to file
    with open(file_path, 'w', encoding='utf-8') as f:
        f.write(csv_content)
    
    return file_path


# =============================================================================
# EXCEL EXPORT FUNCTIONS
# =============================================================================

def generate_users_excel(users: List[User]) -> bytes:
    """
    Generate Excel content for users.
    
    Args:
        users: List of User objects
        
    Returns:
        Excel file content as bytes
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        # Fallback to CSV if openpyxl not installed
        return generate_users_csv(users).encode('utf-8')
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    
    # Header with styling
    headers = ['ID', 'Telegram ID', 'Name', 'Role', 'Class ID', 'Phone', 'Birthday', 'Language', 'Created At']
    ws.append(headers)
    
    # Style header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
    
    # Data rows
    role_map = {1: 'Student', 2: 'Teacher', 3: 'Leader', 4: 'Manager', 5: 'Developer'}
    
    for user in users:
        ws.append([
            user.id,
            user.telegram_id,
            user.name,
            role_map.get(user.role, 'Unknown'),
            user.class_id or '',
            user.phone or '',
            user.birthday.strftime('%Y-%m-%d') if user.birthday else '',
            user.language or 'ar',
            user.created_at.strftime('%Y-%m-%d %H:%M:%S') if user.created_at else ''
        ])
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save to bytes
    from io import BytesIO
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_attendance_excel(records: Optional[List[tuple]] = None) -> bytes:
    """
    Generate Excel content for attendance records.
    
    Args:
        records: Optional list of attendance records (if None, fetches all)
        
    Returns:
        Excel file content as bytes
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return generate_attendance_csv(records).encode('utf-8')
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    
    # Header
    headers = ['ID', 'Student Name', 'Student ID', 'Class ID', 'Date', 'Status', 'Note', 'Marked By', 'Marked At']
    ws.append(headers)
    
    # Style header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
    
    # Fetch records if not provided
    if records is None:
        with get_db() as db:
            records = db.query(
                Attendance,
                User.name.label('student_name'),
                User.telegram_id.label('student_telegram_id')
            ).join(
                User, Attendance.user_id == User.id
            ).all()
    
    # Data rows
    for record in records:
        if hasattr(record, '_asdict'):
            att = record.Attendance if hasattr(record, 'Attendance') else record
            student_name = record.student_name if hasattr(record, 'student_name') else 'Unknown'
            student_tid = record.student_telegram_id if hasattr(record, 'student_telegram_id') else ''
        else:
            att = record
            student_name = att.user.name if att.user else 'Unknown'
            student_tid = att.user.telegram_id if att.user else ''
        
        # Get marker name
        with get_db() as db:
            marker = db.query(User).filter(User.id == att.marked_by).first()
            marker_name = marker.name if marker else 'Unknown'
        
        ws.append([
            att.id,
            student_name,
            student_tid,
            att.class_id or '',
            att.date.strftime('%Y-%m-%d') if isinstance(att.date, date) else str(att.date),
            'Present' if att.status else 'Absent',
            att.note or '',
            marker_name,
            att.created_at.strftime('%Y-%m-%d %H:%M:%S') if att.created_at else ''
        ])
    
    # Auto-adjust columns
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    from io import BytesIO
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def save_excel_to_file(excel_content: bytes, filename: str) -> str:
    """
    Save Excel content to a file.
    
    Args:
        excel_content: Excel content as bytes
        filename: Filename (without path or extension)
        
    Returns:
        Full file path
    """
    import os
    
    # Create exports directory if it doesn't exist
    export_dir = "/home/Bisho/Telegram/exports"
    if not os.path.exists(export_dir):
        os.makedirs(export_dir)
    
    # Generate filename with timestamp
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    full_filename = f"{filename}_{timestamp}.xlsx"
    file_path = os.path.join(export_dir, full_filename)
    
    # Write to file
    with open(file_path, 'wb') as f:
        f.write(excel_content)
    
    return file_path


def generate_class_stats_excel() -> bytes:
    """Generate Excel content for class statistics."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return generate_class_stats_csv().encode('utf-8')
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Class Statistics"
    
    headers = ['Class ID', 'Total Students', 'Total Absences', 'Absences with Reason', 
               'Reason Rate (%)', 'Top Reason 1', 'Top Reason 2', 'Top Reason 3']
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
    
    with get_db() as db:
        classes = db.query(Class).all()
    
    for class_obj in classes:
        users = get_users_by_class(class_obj.id)
        students = [u for u in users if u.role == 1]
        stats = get_attendance_stats_by_class(class_obj.id)
        
        total_absent = stats.get('total_absent', 0)
        with_reason = stats.get('total_with_reason', 0)
        reason_rate = (with_reason / total_absent * 100) if total_absent > 0 else 0
        
        reasons = stats.get('reason_breakdown', {})
        sorted_reasons = sorted(reasons.items(), key=lambda x: x[1], reverse=True)[:3]
        top_reasons = [r[0] for r in sorted_reasons] + ['', '', '']
        
        ws.append([class_obj.id, len(students), total_absent, with_reason, f'{reason_rate:.1f}',
                   top_reasons[0], top_reasons[1], top_reasons[2]])
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    from io import BytesIO
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_full_report_excel() -> bytes:
    """Generate Excel content for comprehensive report."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return generate_full_report_csv().encode('utf-8')
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Full Report"
    
    headers = ['Student ID', 'Student Name', 'Telegram ID', 'Class ID', 'Phone',
               'Total Records', 'Present', 'Absent', 'Rate (%)', 'Latest Date', 'Latest Status']
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
    
    students = get_users_by_role(1)
    
    for student in students:
        with get_db() as db:
            records = db.query(Attendance).filter(Attendance.user_id == student.id).all()
            total_records = len(records)
            present_count = sum(1 for r in records if r.status)
            absent_count = total_records - present_count
            attendance_rate = (present_count / total_records * 100) if total_records > 0 else 0
            
            latest = db.query(Attendance).filter(
                Attendance.user_id == student.id
            ).order_by(Attendance.date.desc()).first()
            
            latest_date = latest.date.strftime('%Y-%m-%d') if latest else ''
            latest_status = 'Present' if (latest and latest.status) else ('Absent' if latest else '')
        
        ws.append([student.id, student.name, student.telegram_id, student.class_id or '',
                   student.phone or '', total_records, present_count, absent_count,
                   f'{attendance_rate:.1f}', latest_date, latest_status])
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    from io import BytesIO
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
